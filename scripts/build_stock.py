import json, re
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "everything-wool-data-spreadsheet.xlsx"
OUT = ROOT / "data/products.json"

def norm(x):
    return re.sub(r"[^a-z0-9]+", "", str(x or "").lower())

if not XLSX.exists():
    raise SystemExit("Missing everything-wool-data-spreadsheet.xlsx")

wb = load_workbook(XLSX, data_only=True)
ws = wb.active
rows = list(ws.iter_rows(values_only=True))
if not rows:
    raise SystemExit("Spreadsheet is empty")

headers = [norm(x) for x in rows[0]]
aliases = {
    "name": ["name","product","productname","title"],
    "category": ["category","type"],
    "price": ["price","saleprice","unitprice"],
    "stock": ["stock","quantity","qty","inventory"],
    "image": ["image","imageurl","photo","photourl"],
    "featured": ["featured","homepage"],
    "description": ["description","details"]
}
idx = {}
for field, names in aliases.items():
    for name in names:
        if name in headers:
            idx[field] = headers.index(name)
            break

if "name" not in idx:
    raise SystemExit("Add a product/name column to the spreadsheet.")

products=[]
for row in rows[1:]:
    if not any(v not in (None,"") for v in row): continue
    def val(field, default=""):
        i=idx.get(field)
        return row[i] if i is not None and i < len(row) else default
    name=val("name")
    if not name: continue
    stock=val("stock",0)
    try: stock=int(float(stock))
    except: stock=0
    price=val("price","")
    try: price=float(price) if price not in ("",None) else ""
    except: pass
    featured=val("featured",False)
    featured=str(featured).strip().lower() in ("true","yes","y","1")
    products.append({
        "name":str(name), "category":str(val("category","")),
        "price":price, "stock":stock, "image":str(val("image","")),
        "featured":featured, "description":str(val("description",""))
    })

OUT.write_text(json.dumps(products, indent=2, ensure_ascii=False), encoding="utf-8")
print(f"Wrote {len(products)} products to {OUT}")
